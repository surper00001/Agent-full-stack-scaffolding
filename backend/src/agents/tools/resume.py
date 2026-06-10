"""简历与求职信生成工具 — 沙箱中 python-docx 排版 .docx 文件。

提供两个工具：
- generate_resume_docx:  生成 6 种排版风格的简历
- generate_cover_letter_docx: 生成配套求职信
"""

from __future__ import annotations

import json
import uuid
from pathlib import Path

from langchain_core.tools import tool
from loguru import logger

from src.core.config import get_settings

_OUTPUT_DIR = Path(get_settings().file_output_dir)
_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ══════════════════════════════════════════════════════════════════════
# 沙箱内 python-docx 排版模板
# ══════════════════════════════════════════════════════════════════════

_RESUME_TEMPLATE = r'''
import json as _json, sys as _sys
try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx not installed")
    _sys.exit(1)

# ── 6 种风格配置 ──
_STYLES = {
    "modern": {
        "font_main":"Calibri","font_heading":"Calibri","font_east":"Microsoft YaHei",
        "c_primary":"1A1A1A","c_accent":"2563EB","c_divider":"2563EB","c_muted":"666666",
        "sz_name":28,"sz_title":13,"sz_section":13,"sz_body":10.5,"sz_contact":9.5,
        "page_margin":0.8,"uppercase":True,
    },
    "classic": {
        "font_main":"Times New Roman","font_heading":"Times New Roman","font_east":"宋体",
        "c_primary":"1A1A1A","c_accent":"2B579A","c_divider":"2B579A","c_muted":"666666",
        "sz_name":26,"sz_title":14,"sz_section":14,"sz_body":11,"sz_contact":10,
        "page_margin":1.0,"uppercase":False,
    },
    "creative": {
        "font_main":"Segoe UI","font_heading":"Segoe UI","font_east":"Microsoft YaHei",
        "c_primary":"FFFFFF","c_accent":"F97316","c_divider":"F97316","c_muted":"AAAAAA",
        "sz_name":24,"sz_title":13,"sz_section":13,"sz_body":10.5,"sz_contact":9.5,
        "page_margin":0.7,"uppercase":True,
    },
    "ats_safe": {
        "font_main":"Arial","font_heading":"Arial","font_east":"Microsoft YaHei",
        "c_primary":"000000","c_accent":"000000","c_divider":"000000","c_muted":"333333",
        "sz_name":20,"sz_title":12,"sz_section":12,"sz_body":11,"sz_contact":10,
        "page_margin":1.0,"uppercase":True,
    },
    "executive": {
        "font_main":"Georgia","font_heading":"Georgia","font_east":"Microsoft YaHei",
        "c_primary":"1A1A1A","c_accent":"1A3A5C","c_divider":"1A3A5C","c_muted":"555555",
        "sz_name":24,"sz_title":12,"sz_section":12,"sz_body":10.5,"sz_contact":9.5,
        "page_margin":0.9,"uppercase":True,
    },
    "academic": {
        "font_main":"Times New Roman","font_heading":"Times New Roman","font_east":"宋体",
        "c_primary":"000000","c_accent":"333333","c_divider":"333333","c_muted":"555555",
        "sz_name":18,"sz_title":12,"sz_section":12,"sz_body":11,"sz_contact":10,
        "page_margin":1.0,"uppercase":False,
    },
}

def _hex(v):
    v=v.lstrip("#");return RGBColor(*[int(v[i:i+2],16) for i in(0,2,4)])

def _rFonts(rPr, s):
    rf=OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"),s["font_main"]);rf.set(qn("w:hAnsi"),s["font_main"])
    rf.set(qn("w:eastAsia"),s["font_east"]);rPr.insert(0,rf)

def _setup(doc,st):
    s=_STYLES[st]
    for sec in doc.sections:
        sec.top_margin=Inches(s["page_margin"])
        sec.bottom_margin=Inches(s["page_margin"])
        sec.left_margin=Inches(s["page_margin"])
        sec.right_margin=Inches(s["page_margin"])
    n=doc.styles["Normal"];f=n.font
    f.name=s["font_main"];f.size=Pt(s["sz_body"]);f.color.rgb=_hex(s["c_primary"])
    _rFonts(n.element.get_or_add_rPr(),s)

def _hdr(doc,name,title,st):
    s=_STYLES[st];p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(name);r.bold=True;r.font.size=Pt(s["sz_name"])
    r.font.color.rgb=_hex(s["c_primary"]);r.font.name=s["font_heading"]
    if title:
        p2=doc.add_paragraph();p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r2=p2.add_run(title);r2.font.size=Pt(s["sz_title"])
        r2.font.color.rgb=_hex(s["c_accent"]);r2.font.name=s["font_main"]
        p2.paragraph_format.space_after=Pt(4)
    _div(doc,st)

def _contact(doc,data,st):
    s=_STYLES[st];parts=[]
    if data.get("phone"):parts.append(data["phone"])
    if data.get("email"):parts.append(data["email"])
    if data.get("city"):parts.append(data["city"])
    if data.get("linkedin"):parts.append(data["linkedin"])
    if data.get("github"):parts.append(data["github"])
    if not parts:return
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("  |  ".join(parts))
    r.font.size=Pt(s["sz_contact"]);r.font.color.rgb=_hex(s["c_primary"]);r.font.name=s["font_main"]
    p.paragraph_format.space_after=Pt(4);_div(doc,st)

def _sect(doc,title,st):
    s=_STYLES[st];p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(3)
    t=title.upper() if s["uppercase"] else title
    r=p.add_run(t);r.bold=True;r.font.size=Pt(s["sz_section"])
    r.font.color.rgb=_hex(s["c_accent"]);r.font.name=s["font_heading"]
    _border(p,s["c_divider"])

def _pg(doc,text,st,indent=False):
    s=_STYLES[st];p=doc.add_paragraph()
    r=p.add_run(text);r.font.size=Pt(s["sz_body"]);r.font.name=s["font_main"]
    if indent:p.paragraph_format.first_line_indent=Cm(0.74)
    p.paragraph_format.space_after=Pt(3)

def _exp(doc,exp,st):
    s=_STYLES[st];company=exp.get("company","");title=exp.get("title","")
    start=exp.get("start_date","");end=exp.get("end_date","")
    dr=f"{start} - {end}" if start else ""
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(2)
    rt=p.add_run(title);rt.bold=True;rt.font.size=Pt(s["sz_body"]);rt.font.name=s["font_heading"]
    rm=p.add_run(f"  |  {company}");rm.font.size=Pt(s["sz_body"]);rm.font.name=s["font_main"]
    if dr:
        rd=p.add_run(f"  |  {dr}");rd.font.size=Pt(s["sz_contact"])
        rd.font.color.rgb=_hex(s["c_muted"]);rd.font.name=s["font_main"]
    for b in exp.get("bullets",[]):_bul(doc,b,st)

def _edu(doc,edu,st):
    s=_STYLES[st];school=edu.get("school","");degree=edu.get("degree","")
    start=edu.get("start_date","");end=edu.get("end_date","");note=edu.get("note","")
    dr=f"{start} - {end}" if start else ""
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(f"{degree}  |  {school}");r.bold=True
    r.font.size=Pt(s["sz_body"]);r.font.name=s["font_heading"]
    if dr:
        rd=p.add_run(f"  |  {dr}");rd.font.size=Pt(s["sz_contact"])
        rd.font.color.rgb=_hex(s["c_muted"]);rd.font.name=s["font_main"]
    if note:_bul(doc,note,st)

def _proj(doc,proj,st):
    s=_STYLES[st];name=proj.get("name","");role=proj.get("role","")
    start=proj.get("start_date","");end=proj.get("end_date","");desc=proj.get("description","")
    dr=f"{start} - {end}" if start else ""
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(name);r.bold=True;r.font.size=Pt(s["sz_body"]);r.font.name=s["font_heading"]
    if role:
        rr=p.add_run(f"  |  {role}");rr.font.size=Pt(s["sz_body"]);rr.font.name=s["font_main"]
    if dr:
        rd=p.add_run(f"  |  {dr}");rd.font.size=Pt(s["sz_contact"])
        rd.font.color.rgb=_hex(s["c_muted"]);rd.font.name=s["font_main"]
    if desc:_bul(doc,desc,st)

def _pub(doc,pub,st):
    """出版物条目：作者 (年份). 标题. 期刊/会议, 卷(期), 页码."""
    s=_STYLES[st]
    text=pub if isinstance(pub,str) else pub.get("citation",str(pub))
    _bul(doc,text,st)

def _skills(doc,skills,st):
    s=_STYLES[st];p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    for i,sk in enumerate(skills):
        r=p.add_run(sk);r.font.size=Pt(s["sz_body"]);r.font.name=s["font_main"]
        r.font.color.rgb=_hex(s["c_accent"])
        if i<len(skills)-1:
            sep=p.add_run("  ·  ");sep.font.size=Pt(8);sep.font.color.rgb=_hex("999999")

def _bul(doc,text,st):
    s=_STYLES[st];p=doc.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.5);p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.space_before=Pt(0)
    r=p.add_run(f"• {text}");r.font.size=Pt(s["sz_body"]);r.font.name=s["font_main"]

def _div(doc,st):
    s=_STYLES[st];p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
    pPr=p._p.get_or_add_pPr();pBdr=OxmlElement("w:pBdr");bt=OxmlElement("w:bottom")
    bt.set(qn("w:val"),"single");bt.set(qn("w:sz"),"4");bt.set(qn("w:space"),"1")
    bt.set(qn("w:color"),s["c_divider"]);pBdr.append(bt);pPr.append(pBdr)

def _border(paragraph,color):
    pPr=paragraph._p.get_or_add_pPr();pBdr=OxmlElement("w:pBdr");bt=OxmlElement("w:bottom")
    bt.set(qn("w:val"),"single");bt.set(qn("w:sz"),"4");bt.set(qn("w:space"),"1")
    bt.set(qn("w:color"),color);pBdr.append(bt);pPr.append(pBdr)

# ── 主执行函数 ──
def execute(data):
    st=data.get("style","modern")
    if st not in _STYLES:st="modern"
    doc=Document();_setup(doc,st)
    _hdr(doc,data.get("full_name",""),data.get("title",""),st)
    _contact(doc,data,st)

    # 摘要
    summary=data.get("summary","")
    if summary:_sect(doc,"Professional Summary" if st!="classic" else "个人总结",st);_pg(doc,summary,st,indent=True)

    # 工作经历
    exps=data.get("experiences",[])
    if exps:_sect(doc,"Experience" if st!="classic" else "工作经历",st)
    for e in exps:_exp(doc,e,st)

    # 教育
    edu=data.get("education",[])
    if edu:_sect(doc,"Education" if st!="classic" else "教育背景",st)
    for e in edu:_edu(doc,e,st)

    # 项目
    projs=data.get("projects",[])
    if projs:_sect(doc,"Projects" if st!="classic" else "项目经验",st)
    for p in projs:_proj(doc,p,st)

    # 出版物 (学术风)
    pubs=data.get("publications",[])
    if pubs:_sect(doc,"Publications" if st!="classic" else "发表论文",st)
    for p in pubs:_pub(doc,p,st)

    # 专利
    patents=data.get("patents",[])
    if patents:_sect(doc,"Patents" if st!="classic" else "专利",st)
    for p in patents:_bul(doc,p if isinstance(p,str) else p.get("title",str(p)),st)

    # 行业演讲 / 董事会 (高管风)
    speaking=data.get("speaking_engagements",[])
    if speaking:_sect(doc,"Speaking & Thought Leadership" if st!="classic" else "行业演讲",st)
    for s in speaking:_bul(doc,s if isinstance(s,str) else s.get("event",str(s)),st)

    board=data.get("board_experience",[])
    if board:_sect(doc,"Board & Advisory" if st!="classic" else "董事会与顾问",st)
    for b in board:_bul(doc,b if isinstance(b,str) else f"{b.get('org','')} — {b.get('role','')}",st)

    # 教学经验 (学术风)
    teaching=data.get("teaching_experience",[])
    if teaching:_sect(doc,"Teaching Experience" if st!="classic" else "教学经验",st)
    for t in teaching:_bul(doc,t if isinstance(t,str) else f"{t.get('course','')} ({t.get('term','')})",st)

    # 基金 / 资助 (学术风)
    grants=data.get("grants",[])
    if grants:_sect(doc,"Grants & Fellowships" if st!="classic" else "基金与资助",st)
    for g in grants:_bul(doc,g if isinstance(g,str) else g.get("title",str(g)),st)

    # 技能
    skills=data.get("skills",[])
    if skills:_sect(doc,"Skills" if st!="classic" else "专业技能",st);_skills(doc,skills,st)

    # 证书 + 语言
    certs=data.get("certifications",[]);langs=data.get("languages",[])
    if certs or langs:
        _sect(doc,"Certifications & Languages" if st!="classic" else "证书与语言",st)
        for c in certs:_bul(doc,c,st)
        for l in langs:_bul(doc,l,st)

    # 志愿 / 荣誉
    awards=data.get("awards",[])
    if awards:_sect(doc,"Awards & Honors" if st!="classic" else "奖项与荣誉",st)
    for a in awards:_bul(doc,a,st)

    volunteer=data.get("volunteer",[])
    if volunteer:_sect(doc,"Volunteer & Community" if st!="classic" else "志愿与社区",st)
    for v in volunteer:_bul(doc,v,st)

    out=data.get("_output_path","/tmp/resume.docx");doc.save(out)
    return {"status":"ok","path":out,"style":st}

if __name__=="__sandbox__":
    print(_json.dumps(execute(_DATA_),ensure_ascii=False))
'''

# ── 求职信模板 ──
_COVER_LETTER_TEMPLATE = r'''
import json as _json, sys as _sys, datetime as _dt
try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx not installed")
    _sys.exit(1)

def execute(data):
    doc=Document()
    # page setup
    for sec in doc.sections:
        sec.top_margin=Inches(1.0);sec.bottom_margin=Inches(1.0)
        sec.left_margin=Inches(1.0);sec.right_margin=Inches(1.0)
    style=doc.styles["Normal"];style.font.name="Calibri";style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.15

    # date
    today=_dt.date.today().strftime("%B %d, %Y")
    p=doc.add_paragraph(today);p.paragraph_format.space_after=Pt(12)

    # recipient
    recipient=data.get("hiring_manager","Hiring Manager")
    company=data.get("company_name","")
    p=doc.add_paragraph();r=p.add_run(recipient)
    if company:
        p2=doc.add_paragraph();p2.add_run(company)
    doc.add_paragraph("")

    # salutation
    salutation=f"Dear {recipient},"
    doc.add_paragraph(salutation).paragraph_format.space_after=Pt(6)

    # body paragraphs
    body=data.get("body_paragraphs",[])
    if not body:
        body=["[Your cover letter content here]"]
    for bp in body:
        p=doc.add_paragraph(bp)
        p.paragraph_format.space_after=Pt(6)

    # close
    doc.add_paragraph("")
    doc.add_paragraph("Sincerely,")
    full_name=data.get("full_name","")
    p=doc.add_paragraph()
    r=p.add_run(full_name);r.bold=True
    if data.get("phone"):
        doc.add_paragraph(data["phone"])
    if data.get("email"):
        doc.add_paragraph(data["email"])

    out=data.get("_output_path","/tmp/cover_letter.docx");doc.save(out)
    return {"status":"ok","path":out}

if __name__=="__sandbox__":
    print(_json.dumps(execute(_DATA_),ensure_ascii=False))
'''

VALID_STYLES = {"modern", "classic", "creative", "ats_safe", "executive", "academic"}


# ══════════════════════════════════════════════════════════════════════
# Tool 1: 简历生成
# ══════════════════════════════════════════════════════════════════════

@tool
async def generate_resume_docx(resume_json: str, style: str = "modern") -> str:
    """
    在沙箱中生成排版精美的简历 .docx 文件（Word 文档），返回下载链接。

    结构化简历数据在沙箱中用 python-docx 引擎排版，支持 6 种专业风格。

    Args:
        resume_json: 简历数据 JSON 字符串，包含：
            full_name, title (必填)
            phone, email, city, linkedin, github (可选联系方式)
            summary (个人总结，3-5句)
            experiences: [{company, title, start_date, end_date, bullets: [str, ...]}]
            education: [{school, degree, start_date, end_date, note}]
            projects: [{name, role, start_date, end_date, description}]
            skills: [str, ...]
            certifications: [str, ...]
            languages: [str, ...]
            publications: [str|citation, ...]      (学术风)
            patents: [str, ...]                     (学术/高管风)
            speaking_engagements: [str, ...]        (高管风)
            board_experience: [{org, role}]         (高管风)
            teaching_experience: [{course, term}]   (学术风)
            grants: [{title, ...}]                  (学术风)
            awards: [str, ...]
            volunteer: [str, ...]
        style: 排版风格 — modern/classic/creative/ats_safe/executive/academic
            - modern (推荐): 无衬线现代风 — 互联网/外企/SaaS
            - classic: 宋体传统 — 国企/金融/法律
            - creative: 创意设计 — 设计师/媒体
            - ats_safe: 极致ATS兼容 — 海投/大厂网申
            - executive: 高管风 — VP/合伙人/董事会
            - academic: 学术CV — 教职/博后/研究员

    Returns:
        JSON: {filename, download_url, mime_type, size_bytes, style, message}
    """
    try:
        data = json.loads(resume_json)
    except json.JSONDecodeError as e:
        return json.dumps({"error": f"JSON parse error: {e}", "hint": "Please provide valid JSON"}, ensure_ascii=False)

    if not data.get("full_name"):
        return json.dumps({"error": "full_name is required"}, ensure_ascii=False)

    if style not in VALID_STYLES:
        style = "modern"

    file_id = uuid.uuid4().hex[:12]
    safe_name = "".join(c for c in data.get("full_name", "resume") if c.isalnum() or c == "_")[:20]
    filename = f"{safe_name}_{file_id}.docx"
    output_path = _OUTPUT_DIR / filename
    data["_output_path"] = str(output_path)
    data["style"] = style

    try:
        from src.harness.sandbox.manager import get_sandbox_manager
        logger.info(f"Sandbox resume generation: {filename} (style={style})")
        manager = get_sandbox_manager()
        code = f"_DATA_ = {json.dumps(data, ensure_ascii=False)}\n{_RESUME_TEMPLATE}"
        result = await manager.execute_code(code, timeout=30)
        if not result.success:
            logger.error(f"Sandbox error: {result.stderr or result.stdout}")
            return json.dumps({"error": "Resume generation failed in sandbox", "detail": result.stderr or result.stdout}, ensure_ascii=False)
    except Exception as e:
        logger.exception(f"Sandbox init failed: {e}")
        return json.dumps({"error": f"Sandbox unavailable: {e}"}, ensure_ascii=False)

    if not output_path.exists():
        return json.dumps({"error": "Output file not found after sandbox execution"}, ensure_ascii=False)

    settings = get_settings()
    download_url = f"{settings.file_download_url_prefix}/{filename}"
    size_kb = output_path.stat().st_size / 1024

    style_labels = {"modern":"Modern","classic":"Classic","creative":"Creative","ats_safe":"ATS-Safe","executive":"Executive","academic":"Academic"}
    return json.dumps({
        "filename": filename,
        "download_url": download_url,
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": output_path.stat().st_size,
        "size_display": f"{size_kb:.1f} KB",
        "style": style,
        "style_label": style_labels.get(style, style),
        "message": (
            f"✅ Resume generated!\n"
            f"\U0001f4c4 File: {filename}\n"
            f"\U0001f4cf Size: {size_kb:.1f} KB\n"
            f"\U0001f3a8 Style: {style_labels.get(style, style)}\n"
            f"\U0001f4e5 Download and edit in Microsoft Word\n"
            f"\U0001f4a1 Reply to switch style ({'/'.join(sorted(VALID_STYLES))}) or update content"
        ),
    }, ensure_ascii=False)


# ══════════════════════════════════════════════════════════════════════
# Tool 2: 求职信生成
# ══════════════════════════════════════════════════════════════════════

@tool
async def generate_cover_letter_docx(
    cover_letter_json: str,
) -> str:
    """
    在沙箱中生成求职信 (Cover Letter) .docx 文件，返回下载链接。

    Args:
        cover_letter_json: JSON 字符串，包含：
            full_name, phone?, email? (发信人信息)
            hiring_manager (收信人姓名，默认 "Hiring Manager")
            company_name (目标公司)
            body_paragraphs: [str, ...] (3-5 段的正文，Agent 按四段式结构撰写)
                - Para 1: Hook — 你是谁 + 为什么对这个职位感兴趣
                - Para 2: Proof — 2-3 个最相关成就，STAR 展开
                - Para 3: Connection — 对公司/产品的了解和真诚兴趣
                - Para 4: Close — 面谈意愿 + 感谢

    Returns:
        JSON: {filename, download_url, mime_type, size_bytes, message}
    """
    try:
        data = json.loads(cover_letter_json)
    except json.JSONDecodeError as e:
        return json.dumps({"error": f"JSON parse error: {e}"}, ensure_ascii=False)

    if not data.get("full_name"):
        return json.dumps({"error": "full_name is required"}, ensure_ascii=False)
    if not data.get("body_paragraphs"):
        return json.dumps({"error": "body_paragraphs is required (3-5 paragraphs)"}, ensure_ascii=False)

    file_id = uuid.uuid4().hex[:12]
    safe_name = "".join(c for c in data.get("full_name", "cover_letter") if c.isalnum() or c == "_")[:20]
    filename = f"cover_letter_{safe_name}_{file_id}.docx"
    output_path = _OUTPUT_DIR / filename
    data["_output_path"] = str(output_path)

    try:
        from src.harness.sandbox.manager import get_sandbox_manager
        logger.info(f"Sandbox cover letter: {filename}")
        manager = get_sandbox_manager()
        code = f"_DATA_ = {json.dumps(data, ensure_ascii=False)}\n{_COVER_LETTER_TEMPLATE}"
        result = await manager.execute_code(code, timeout=20)
        if not result.success:
            return json.dumps({"error": "Cover letter generation failed", "detail": result.stderr or result.stdout}, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"error": f"Sandbox unavailable: {e}"}, ensure_ascii=False)

    if not output_path.exists():
        return json.dumps({"error": "Output file not found"}, ensure_ascii=False)

    settings = get_settings()
    size_kb = output_path.stat().st_size / 1024
    return json.dumps({
        "filename": filename,
        "download_url": f"{settings.file_download_url_prefix}/{filename}",
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": output_path.stat().st_size,
        "size_display": f"{size_kb:.1f} KB",
        "message": f"✅ Cover letter generated!\n\U0001f4c4 {filename} ({size_kb:.1f} KB)\n\U0001f4e5 Download and edit in Word",
    }, ensure_ascii=False)
